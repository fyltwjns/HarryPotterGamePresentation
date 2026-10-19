import pptx
import os
import hashlib

prs = pptx.Presentation('presentation/LSGTDA.pptx')

asset_hashes = {}
for fname in os.listdir('assets'):
    fpath = os.path.join('assets', fname)
    if os.path.isfile(fpath):
        with open(fpath, 'rb') as f:
            asset_hashes[hashlib.md5(f.read()).hexdigest()] = fname

def get_asset(blob):
    h = hashlib.md5(blob).hexdigest()
    return asset_hashes.get(h, 'unmatched')

# Audit each screen:
# 1: Title
# 2: Letter
# 3: Rules
# 4: Grid
# 5: Question template
# 20: Supplies
# 21: Friend
# 22-24: Letter
# 25: Duel
# 26: Test luck 1
# 27-29: Rival Duel
# 30-32: Voldemort
# 33: 3 Classes
# 34: Owl
# 35: Test luck 2
# 36: Bowls
# 37-39: Hippogriff
# 40: Snitch
# 41: Sorting Hats
# 42: Patronus

slides_to_check = [1, 2, 3, 4, 5, 20, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30, 31, 32, 33, 34, 35, 36, 37, 38, 39, 40, 41, 42]

print("=== SLIDE AUDIT DETAILS ===")
for num in slides_to_check:
    slide = prs.slides[num - 1]
    # Background
    bg_asset = 'None'
    try:
        blip = slide.background._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
        if blip is not None:
            rId = blip.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            rel = slide.part.rels.get(rId)
            if rel:
                bg_asset = get_asset(rel.target_part.blob)
    except Exception as e:
        pass
    
    shapes_info = []
    for s in slide.shapes:
        txt = s.text_frame.text.strip().replace('\n', ' ') if s.has_text_frame else ''
        img = get_asset(s.image.blob) if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE else ''
        x, y, w, h = s.left//10000, s.top//10000, s.width//10000, s.height//10000
        shapes_info.append(f"{s.name[:18]}: pos=({x},{y},{w},{h}) img={img} txt={repr(txt[:40])}")
    
    print(f"\n--- Slide {num} (BG: {bg_asset}) ---")
    for si in shapes_info[:6]:
        print("  " + si)
    if len(shapes_info) > 6:
        print(f"  ... and {len(shapes_info) - 6} more shapes")
