import pptx
import json

prs = pptx.Presentation('presentation/LSGTDA.pptx')
slide = prs.slides[32]  # Slide 33

print(f"=== SLIDE 33 SHAPES ({len(slide.shapes)}) ===")
shapes_info = []
for i, s in enumerate(slide.shapes):
    text = ""
    if s.has_text_frame:
        text = "\n".join([p.text for p in s.text_frame.paragraphs if p.text.strip()])
    
    img_name = ""
    if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
        # check image part
        try:
            img_name = s.image.filename
        except:
            img_name = "pic"
            
    info = {
        "idx": i,
        "id": s.shape_id,
        "name": s.name,
        "text": text,
        "img": img_name,
        "left": s.left,
        "top": s.top,
        "width": s.width,
        "height": s.height
    }
    shapes_info.append(info)
    if text or img_name:
        print(f"[{i}] id={s.shape_id} name='{s.name}' text='{text}' img='{img_name}' pos=({s.left},{s.top})")

with open('tools/slide33_shapes.json', 'w', encoding='utf-8') as f:
    json.dump(shapes_info, f, indent=2, ensure_ascii=False)

for s in shapes_info:
    print(f"[{s['idx']}] id={s['id']} name='{s['name']}' text={repr(s['text'])} img='{s['img']}' type={s.get('shape_type', '')} pos=({s['left']},{s['top']}) size=({s['width']},{s['height']})")
