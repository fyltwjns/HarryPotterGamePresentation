import pptx
import os
import hashlib
import json

prs = pptx.Presentation('presentation/LSGTDA.pptx')

asset_hashes = {}
for fname in os.listdir('assets'):
    fpath = os.path.join('assets', fname)
    if os.path.isfile(fpath):
        with open(fpath, 'rb') as f:
            asset_hashes[hashlib.md5(f.read()).hexdigest()] = fname

def get_asset(blob):
    h = hashlib.md5(blob).hexdigest()
    return asset_hashes.get(h, 'unmatched')

audit_data = []

for idx, slide in enumerate(prs.slides):
    slide_num = idx + 1
    bg_asset = 'None'
    try:
        blip = slide.background._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
        if blip is not None:
            rId = blip.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            rel = slide.part.rels.get(rId)
            if rel:
                bg_asset = get_asset(rel.target_part.blob)
    except Exception as e:
        pass

    shapes_list = []
    for s in slide.shapes:
        info = {
            'name': s.name,
            'left': s.left,
            'top': s.top,
            'width': s.width,
            'height': s.height,
            'type': str(s.shape_type),
            'text': '',
            'image': '',
            'font_names': [],
            'font_sizes': []
        }
        if s.has_text_frame:
            info['text'] = s.text_frame.text.strip()
            for p in s.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.name and r.font.name not in info['font_names']:
                        info['font_names'].append(r.font.name)
                    if r.font.size and r.font.size.pt not in info['font_sizes']:
                        info['font_sizes'].append(r.font.size.pt)
        if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            info['image'] = get_asset(s.image.blob)
        shapes_list.append(info)

    audit_data.append({
        'slide_number': slide_num,
        'background': bg_asset,
        'shape_count': len(shapes_list),
        'shapes': shapes_list
    })

with open('tools/full_slide_audit.json', 'w', encoding='utf-8') as f:
    json.dump(audit_data, f, ensure_ascii=False, indent=2)

print(f"Audited {len(audit_data)} slides successfully.")
