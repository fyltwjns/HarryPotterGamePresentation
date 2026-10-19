import pptx
import hashlib
import os

prs = pptx.Presentation('presentation/LSGTDA.pptx')
slide = prs.slides[32]

# Map assets to md5 hashes
asset_hashes = {}
for fname in os.listdir('assets'):
    fpath = os.path.join('assets', fname)
    if os.path.isfile(fpath):
        with open(fpath, 'rb') as f:
            asset_hashes[hashlib.md5(f.read()).hexdigest()] = fname

print("=== SLIDE 33 CLASSES & IMAGES ===")
classes = []
for i in range(7, 17):
    s = slide.shapes[i]
    cls_name = ""
    img_asset = "NONE"
    for sub in s.shapes:
        if sub.has_text_frame and sub.text.strip():
            cls_name = sub.text.strip().replace('\x0b', ' ')
        if sub.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            h = hashlib.md5(sub.image.blob).hexdigest()
            img_asset = asset_hashes.get(h, f"UNKNOWN_{h[:8]}")
    classes.append((s.left, s.top, cls_name, img_asset, s.shape_id))

# Sort by y (row), then by x (col)
classes.sort(key=lambda c: (c[1] // 1000000, c[0]))
print("\n=== SLIDE 33 CLASSES SORTED BY GRID POSITION ===")
for idx, c in enumerate(classes):
    print(f"[{idx}] pos=({c[0]}, {c[1]}) -> Class: '{c[2]}', Book: '{c[3]}', ShapeID: {c[4]}")
