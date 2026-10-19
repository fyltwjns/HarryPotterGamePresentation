import pptx

prs = pptx.Presentation('presentation/LSGTDA.pptx')
slide = prs.slides[32]

print("=== ALL SHAPES IN SLIDE 33 ===")
for i, s in enumerate(slide.shapes):
    text = ""
    if s.has_text_frame:
        text = s.text.strip().replace('\n', ' ')
    elif s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
        text = " [GROUP: " + ", ".join([sub.text.strip().replace('\x0b', ' ') for sub in s.shapes if sub.has_text_frame and sub.text.strip()]) + "]"
    print(f"[{i:2d}] id={s.shape_id:3d} name={s.name:<25} pos=({s.left:7d}, {s.top:7d}) size=({s.width:7d}, {s.height:7d}) text={text}")
